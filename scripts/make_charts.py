"""Charts for the EdTech Feature Prioritization and Roadmap project (hypothetical data).

RICE = Reach x Impact x Confidence / Effort, recomputed from the Feature_Backlog inputs.
Usage: python make_charts.py <workbook.xlsx> <out_dir>
"""
import re
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd

SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)

# Reference palette (dataviz skill). MoSCoW is ordered, so it uses an ordinal blue ramp
# (steps 650 / 450 / 250, none lighter than 250) with a neutral gray for "Won't Have".
SURFACE, INK, INK2, MUTED = "#fcfcfb", "#0b0b0b", "#52514e", "#898781"
GRID, AXIS, BLUE = "#e1e0d9", "#c3c2b7", "#2a78d6"
MOSCOW = {"Must Have": "#104281", "Should Have": "#2a78d6", "Could Have": "#86b6ef", "Won't Have": "#c3c2b7"}

plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 11, "text.color": INK,
    "axes.edgecolor": AXIS, "axes.labelcolor": INK2, "xtick.color": MUTED, "ytick.color": MUTED,
    "axes.facecolor": SURFACE, "figure.facecolor": SURFACE, "savefig.facecolor": SURFACE,
})
NOTE = "Hypothetical data"

wb = openpyxl.load_workbook(SRC, data_only=True)

# ---- backlog
fb = wb["Feature_Backlog"]
hdr = next(r for r in range(1, 10) if fb.cell(r, 1).value == "#")
rows = [[fb.cell(r, c).value for c in range(1, 12)] for r in range(hdr + 1, fb.max_row + 1) if fb.cell(r, 1).value is not None]
bl = pd.DataFrame(rows, columns=["n", "feature", "source", "cat", "reach", "impact", "conf", "effort", "rice_cell", "moscow", "q"])
bl["value"] = bl.reach * bl.impact * bl.conf / 100
bl["rice"] = bl.value / bl.effort

# ---- roadmap (12 initiatives in quarter blocks)
pr = wb["Product_Roadmap"]
road, quarter = [], None
for r in range(1, pr.max_row + 1):
    a = pr.cell(r, 1).value
    if isinstance(a, str) and re.match(r"Q\d —", a):
        quarter = a
    elif quarter and a and a != "Initiative" and pr.cell(r, 6).value:
        road.append((quarter, a, pr.cell(r, 3).value, int(re.match(r"\d+", str(pr.cell(r, 6).value)).group())))
rd = pd.DataFrame(road, columns=["quarter", "feature", "moscow", "weeks"])
bl["in_roadmap"] = bl.feature.isin(rd.feature)
assert bl.in_roadmap.sum() == len(rd) == 12, (bl.in_roadmap.sum(), len(rd))


def heading(fig, title, subtitle, y1=0.975, y2=0.935):
    fig.text(0.025, y1, title, fontsize=14, fontweight="bold", color=INK, va="top")
    fig.text(0.025, y2, subtitle, fontsize=10, color=INK2, va="top")


def clean(ax, grid_axis):
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color(AXIS)
    ax.spines["bottom"].set_color(AXIS)
    ax.grid(axis=grid_axis, color=GRID, linewidth=1)
    ax.set_axisbelow(True)
    ax.tick_params(length=0)


def moscow_legend(ax, cats, loc="lower right"):
    handles = [matplotlib.patches.Patch(color=MOSCOW[c], label=c) for c in cats]
    ax.legend(handles=handles, loc=loc, frameon=False, fontsize=10, labelcolor=INK2)


# 1. RICE ranking for all 25 items, coloured by MoSCoW, roadmap items marked
s = bl.sort_values("rice")
fig, ax = plt.subplots(figsize=(7.4, 8.4))
y = np.arange(len(s))
ax.barh(y, s.rice, height=0.62, color=[MOSCOW[m] for m in s.moscow])
for yi, (v, inr) in enumerate(zip(s.rice, s.in_roadmap)):
    ax.text(v + 10, yi, f"{v:,.0f}", va="center", color=INK, fontsize=9.5, fontweight="bold" if inr else "normal")
ax.set_yticks(y, [("● " if inr else "") + f for f, inr in zip(s.feature, s.in_roadmap)], fontsize=9.5)
for lab, inr in zip(ax.get_yticklabels(), s.in_roadmap):
    lab.set_color(INK if inr else MUTED)
ax.set_xlim(0, s.rice.max() * 1.12)
clean(ax, "x")
ax.spines["bottom"].set_visible(False)
moscow_legend(ax, list(MOSCOW))
heading(fig, "MoSCoW, not RICE alone, decides the roadmap",
        f"RICE score for all 25 backlog items · ● = one of the 12 roadmap initiatives · {NOTE}", 0.985, 0.955)
fig.tight_layout(rect=[0, 0, 1, 0.93])
fig.savefig(OUT / "01_rice_ranking.png", dpi=200)
plt.close(fig)

# 2. Value vs effort (prioritization matrix); roadmap vs not in roadmap
fig, ax = plt.subplots(figsize=(6.8, 4.6))
out = bl[~bl.in_roadmap]
inn = bl[bl.in_roadmap]
ax.scatter(out.effort, out.value, s=70, color="#c3c2b7", edgecolor=SURFACE, linewidth=2, label="Not on roadmap", zorder=3)
ax.scatter(inn.effort, inn.value, s=70, color=BLUE, edgecolor=SURFACE, linewidth=2, label="On roadmap", zorder=4)
label = {"Video Playback Speed Control": (6, -14), "Certificate Generation": (6, 6), "Progress Dashboard": (6, 6),
         "Accessibility Mode (WCAG 2.1)": (6, 6), "Mobile App — iOS & Android": (6, -12),
         "Proctoring Integration": (-6, 8)}
for _, r in bl[bl.feature.isin(label)].iterrows():
    dx, dy = label[r.feature]
    ax.annotate(r.feature.replace(" (WCAG 2.1)", "").replace(" — iOS & Android", "").replace(" (Weekly Summary)", ""),
                (r.effort, r.value), xytext=(dx, dy), textcoords="offset points", fontsize=9, color=INK2,
                ha="right" if dx < 0 else "left")
ax.set_xlabel("Effort (weeks)")
ax.set_ylabel("Value (reach × impact × confidence)")
ax.set_xlim(0, 20)
ax.set_ylim(0, bl.value.max() * 1.15)
ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v/1000:g}K" if v >= 1000 else f"{v:.0f}"))
ax.xaxis.set_major_locator(matplotlib.ticker.MultipleLocator(4))
clean(ax, "both")
ax.legend(loc="upper right", frameon=False, fontsize=10, labelcolor=INK2)
heading(fig, "The top three RICE scores all take four weeks or less",
        f"Each dot is a backlog item; RICE = value ÷ effort · {NOTE}", 0.975, 0.91)
fig.tight_layout(rect=[0, 0, 1, 0.86])
fig.savefig(OUT / "02_value_vs_effort.png", dpi=200)
plt.close(fig)

# 3. Roadmap: effort by initiative within each quarter
fig, ax = plt.subplots(figsize=(7.4, 6.2))
labels, ys, y0 = [], [], 0
qtot = rd.groupby("quarter", sort=False).weeks.sum()
for q, grp in rd.groupby("quarter", sort=False):
    for _, r in grp.iterrows():
        ax.barh(y0, r.weeks, height=0.62, color=MOSCOW[r.moscow])
        ax.text(r.weeks + 0.2, y0, f"{r.weeks} wks", va="center", color=INK, fontsize=9.5)
        labels.append(r.feature); ys.append(y0); y0 -= 1
    ax.text(-0.2, y0 + len(grp) + 0.55, f"{q.split(' (')[0]} · {qtot[q]} weeks", ha="left", va="center",
            fontsize=10.5, fontweight="bold", color=INK, transform=ax.get_yaxis_transform())
    y0 -= 1.0
ax.set_yticks(ys, labels, fontsize=9.5)
ax.set_xlim(0, rd.weeks.max() * 1.18)
ax.set_xlabel("Estimated effort (weeks)")
clean(ax, "x")
ax.spines["bottom"].set_visible(False)
moscow_legend(ax, ["Must Have", "Should Have"], loc="center right")
heading(fig, f"12 initiatives over three quarters, {rd.weeks.sum()} weeks of effort",
        f"Roadmap initiatives by quarter and estimated effort · {NOTE}", 0.98, 0.945)
fig.tight_layout(rect=[0, 0, 1, 0.92])
fig.savefig(OUT / "03_roadmap_by_quarter.png", dpi=200)
plt.close(fig)

# 4. Baseline vs target for initiatives whose KPI is a percentage
sm = wb["Success_Metrics"]
pts = []
for r in range(4, sm.max_row + 1):
    name, kpi, base, tgt = (sm.cell(r, c).value for c in range(1, 5))
    m1 = re.fullmatch(r"(\d+(?:\.\d+)?)%", str(base or "").strip())
    m2 = re.match(r"(\d+(?:\.\d+)?)%", str(tgt or "").strip())
    if name and m1 and m2:
        pts.append((name, kpi, float(m1.group(1)), float(m2.group(1))))
pt = pd.DataFrame(pts, columns=["name", "kpi", "base", "target"])
pt["lift"] = pt.target - pt.base
pt = pt.sort_values("lift")
fig, ax = plt.subplots(figsize=(7.2, 4.6))
y = np.arange(len(pt))
for yi, r in zip(y, pt.itertuples()):
    ax.plot([r.base, r.target], [yi, yi], color=AXIS, linewidth=2, zorder=2)
    ax.text(r.target + 2, yi, f"+{r.lift:.0f} pts", va="center", color=INK, fontsize=10, fontweight="bold")
ax.scatter(pt.base, y, s=70, color="#898781", edgecolor=SURFACE, linewidth=2, zorder=3, label="Baseline")
ax.scatter(pt.target, y, s=70, color=BLUE, edgecolor=SURFACE, linewidth=2, zorder=4, label="Target")
names = [("Overall" if n == "OVERALL TARGET" else n) + f"\n{k}" for n, k in zip(pt.name, pt.kpi)]
ax.set_yticks(y, names, fontsize=9)
ax.set_xlim(0, 100)
ax.xaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(100))
clean(ax, "x")
ax.spines["bottom"].set_visible(False)
ax.legend(loc="lower right", frameon=False, fontsize=10, labelcolor=INK2)
top = pt.iloc[-1]
heading(fig, f"{top['name']} carries the biggest target (+{top.lift:.0f} points)",
        f"Baseline and target for initiatives measured in percentages · Hypothetical targets", 0.975, 0.91)
fig.tight_layout(rect=[0, 0, 1, 0.86])
fig.savefig(OUT / "04_targets_vs_baseline.png", dpi=200)
plt.close(fig)

print("charts written to", OUT)
